from fastapi import FastAPI, HTTPException, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter

import qrcode
import secrets
import hashlib
import re

from datetime import datetime


# =========================================================
# APP
# =========================================================

app = FastAPI(
    title="IEEE CIS Attendance System"
)


app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# PATHS
# =========================================================

BASE_DIR = Path(__file__).resolve().parent

# Main/master Excel file
EXCEL_FILE = BASE_DIR / "attendance.xlsx"

# Permanent QR codes
QR_FOLDER = BASE_DIR / "qr_codes"

QR_FOLDER.mkdir(exist_ok=True)

# Separate Excel files for every meeting
DATA_FOLDER = BASE_DIR / "data"

DATA_FOLDER.mkdir(exist_ok=True)


# =========================================================
# PASSWORD HELPERS
# =========================================================

def hash_password(password: str):

    return hashlib.sha256(
        password.encode()
    ).hexdigest()


def verify_password(
    password: str,
    password_hash: str
):

    return hash_password(password) == password_hash


# =========================================================
# SAFE FILE NAME
# =========================================================

def safe_filename(name: str):

    name = re.sub(
        r'[<>:"/\\|?*]+',
        "_",
        str(name)
    )

    name = re.sub(
        r"\s+",
        "_",
        name.strip()
    )

    if not name:
        name = "Meeting"

    return name[:80]


# =========================================================
# SESSION EXCEL PATH
# =========================================================

def session_excel_path(
    session_id: str,
    event: str
):

    filename = (
        f"{safe_filename(event)}_"
        f"{session_id}.xlsx"
    )

    return DATA_FOLDER / filename


# =========================================================
# CREATE INDIVIDUAL SESSION EXCEL
# =========================================================

def create_session_excel(
    session_id: str,
    event: str
):

    path = session_excel_path(
        session_id,
        event
    )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Attendance"

    sheet.append([
        "Sr. No.",
        "Name",
        "Branch",
        "Section",
        "Roll No.",
        "Email",
        "Date",
        "Time",
        "Status"
    ])

    save_session_excel(
        workbook,
        path
    )

    workbook.close()


# =========================================================
# SAVE SESSION EXCEL WITH COLUMN WIDTH
# =========================================================

def save_session_excel(
    workbook,
    path
):

    for sheet in workbook.worksheets:

        for column_cells in sheet.columns:

            max_length = 0

            for cell in column_cells:

                try:

                    length = len(
                        str(cell.value)
                    )

                    if length > max_length:

                        max_length = length

                except Exception:

                    pass

            column_letter = get_column_letter(
                column_cells[0].column
            )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                40
            )

    workbook.save(path)


# =========================================================
# APPEND ATTENDANCE TO SESSION EXCEL
# =========================================================

def append_to_session_excel(
    session_id: str,
    event: str,
    student: dict,
    now: datetime
):

    path = session_excel_path(
        session_id,
        event
    )

    # Create if missing
    if not path.exists():

        create_session_excel(
            session_id,
            event
        )


    workbook = load_workbook(
        path
    )

    sheet = workbook[
        "Attendance"
    ]


    # Sr. No.
    sr_no = sheet.max_row


    sheet.append([

        sr_no,

        student.get(
            "name",
            ""
        ),

        student.get(
            "branch",
            ""
        ),

        student.get(
            "section",
            ""
        ),

        student.get(
            "roll_no",
            ""
        ),

        student.get(
            "email",
            ""
        ),

        now.strftime(
            "%Y-%m-%d"
        ),

        now.strftime(
            "%H:%M:%S"
        ),

        "Present"

    ])


    save_session_excel(
        workbook,
        path
    )

    workbook.close()


# =========================================================
# EXCEL CREATION
# =========================================================

def create_excel():

    workbook = Workbook()


    # =====================================================
    # USERS
    # =====================================================

    users = workbook.active

    users.title = "Users"

    users.append([
        "User ID",
        "Name",
        "Email",
        "Password Hash",
        "Role",
        "Active",
        "Created Date"
    ])


    # =====================================================
    # STUDENTS
    # =====================================================

    students = workbook.create_sheet(
        "Students"
    )

    students.append([
        "Student ID",
        "Name",
        "Roll No.",
        "Section",
        "Branch",
        "Email",
        "Password Hash",
        "QR Token",
        "Registered Date"
    ])


    # =====================================================
    # SESSIONS
    # =====================================================

    sessions = workbook.create_sheet(
        "Sessions"
    )

    sessions.append([
        "Session ID",
        "Event",
        "Date",
        "Start Time",
        "Created By"
    ])


    # =====================================================
    # ATTENDANCE
    # =====================================================

    attendance = workbook.create_sheet(
        "Attendance"
    )

    attendance.append([
        "Session ID",
        "Event",
        "Student ID",
        "Name",
        "Roll No.",
        "Section",
        "Branch",
        "Email",
        "Date",
        "Time",
        "Status"
    ])


    # =====================================================
    # DEFAULT HEAD
    # =====================================================

    users.append([

        "HEAD001",

        "IEEE CIS Head",

        "head@ieeecis.com",

        hash_password(
            "ieeecis123"
        ),

        "HEAD",

        "YES",

        datetime.now().strftime(
            "%Y-%m-%d"
        )

    ])


    workbook.save(
        EXCEL_FILE
    )


# =========================================================
# GET WORKBOOK
# =========================================================

def get_workbook():

    if not EXCEL_FILE.exists():

        create_excel()


    try:

        workbook = load_workbook(
            EXCEL_FILE
        )


        required_sheets = [

            "Users",

            "Students",

            "Sessions",

            "Attendance"

        ]


        for sheet in required_sheets:

            if sheet not in workbook.sheetnames:

                workbook.close()

                create_excel()

                return load_workbook(
                    EXCEL_FILE
                )


        return workbook


    except (
        InvalidFileException,
        OSError,
        ValueError,
        KeyError
    ):

        backup_name = (
            BASE_DIR /
            f"attendance_corrupt_"
            f"{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        )


        try:

            EXCEL_FILE.rename(
                backup_name
            )

        except Exception:

            pass


        create_excel()

        return load_workbook(
            EXCEL_FILE
        )


# =========================================================
# MODELS
# =========================================================

class LoginRequest(BaseModel):

    user_id: str

    password: str


class StudentCreate(BaseModel):

    student_id: str

    name: str

    roll_no: str = ""

    section: str = ""

    branch: str = ""

    email: str = ""

    password: str


class SessionCreatorCreate(BaseModel):

    user_id: str

    name: str

    email: str

    password: str


class SessionCreate(BaseModel):

    event: str


class AttendanceRequest(BaseModel):

    session_id: str

    qr_token: str


# =========================================================
# FIND USER
# =========================================================

def find_user(
    workbook,
    user_id
):

    sheet = workbook[
        "Users"
    ]


    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == str(user_id):

            return {

                "user_id": row[0],

                "name": row[1],

                "email": row[2],

                "password_hash": row[3],

                "role": row[4],

                "active": row[5]

            }


    return None


# =========================================================
# AUTHENTICATION
# =========================================================

def authenticate_user(
    user_id,
    password
):

    workbook = get_workbook()

    user = find_user(
        workbook,
        user_id
    )

    workbook.close()


    if not user:

        raise HTTPException(
            status_code=401,
            detail="Invalid User ID or Password"
        )


    if user["active"] != "YES":

        raise HTTPException(
            status_code=403,
            detail="This account is inactive"
        )


    if not verify_password(
        password,
        user["password_hash"]
    ):

        raise HTTPException(
            status_code=401,
            detail="Invalid User ID or Password"
        )


    token = secrets.token_urlsafe(
        32
    )


    return {

        "token": token,

        "user_id":
            user["user_id"],

        "name":
            user["name"],

        "role":
            user["role"]

    }


# =========================================================
# ACTIVE LOGIN TOKENS
# =========================================================

ACTIVE_TOKENS = {}


# =========================================================
# GET CURRENT USER
# =========================================================

def get_current_user(
    authorization: str | None
):

    if not authorization:

        raise HTTPException(
            status_code=401,
            detail="Login required"
        )


    if not authorization.startswith(
        "Bearer "
    ):

        raise HTTPException(
            status_code=401,
            detail="Invalid authorization"
        )


    token = authorization.replace(
        "Bearer ",
        ""
    )


    user = ACTIVE_TOKENS.get(
        token
    )


    if not user:

        raise HTTPException(
            status_code=401,
            detail="Invalid or expired login"
        )


    return user


# =========================================================
# LOGIN
# =========================================================

@app.post("/login")
def login(
    data: LoginRequest
):

    result = authenticate_user(
        data.user_id,
        data.password
    )


    ACTIVE_TOKENS[
        result["token"]
    ] = {

        "user_id":
            result["user_id"],

        "name":
            result["name"],

        "role":
            result["role"]

    }


    return result


# =========================================================
# LOGOUT
# =========================================================

@app.post("/logout")
def logout(
    authorization: str | None = Header(None)
):

    if authorization:

        token = authorization.replace(
            "Bearer ",
            ""
        )

        ACTIVE_TOKENS.pop(
            token,
            None
        )


    return {
        "message":
            "Logged out successfully"
    }


# =========================================================
# REGISTER STUDENT
# HEAD ONLY
# =========================================================

@app.post("/students")
def register_student(
    data: StudentCreate,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] != "HEAD":

        raise HTTPException(
            status_code=403,
            detail="Only Head can register students"
        )


    workbook = get_workbook()

    students = workbook[
        "Students"
    ]

    users = workbook[
        "Users"
    ]


    student_id = data.student_id.strip()


    if not student_id:

        workbook.close()

        raise HTTPException(
            status_code=400,
            detail="Student ID is required"
        )


    # Check existing student

    for row in students.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == student_id:

            workbook.close()

            raise HTTPException(
                status_code=400,
                detail="Student ID already exists"
            )


    # Check existing user

    if find_user(
        workbook,
        student_id
    ):

        workbook.close()

        raise HTTPException(
            status_code=400,
            detail="User ID already exists"
        )


    qr_token = secrets.token_urlsafe(
        24
    )


    today = datetime.now().strftime(
        "%Y-%m-%d"
    )


    students.append([

        student_id,

        data.name,

        data.roll_no,

        data.section,

        data.branch,

        data.email,

        hash_password(
            data.password
        ),

        qr_token,

        today

    ])


    users.append([

        student_id,

        data.name,

        data.email,

        hash_password(
            data.password
        ),

        "STUDENT",

        "YES",

        today

    ])


    workbook.save(
        EXCEL_FILE
    )

    workbook.close()


    # Permanent QR

    qr = qrcode.make(
        qr_token
    )


    qr_path = (
        QR_FOLDER /
        f"{student_id}.png"
    )


    qr.save(
        qr_path
    )


    return {

        "message":
            "Student registered successfully",

        "student_id":
            student_id,

        "qr_token":
            qr_token,

        "qr_url":
            f"/qr/{student_id}"

    }


# =========================================================
# GET STUDENTS
# HEAD ONLY
# =========================================================

@app.get("/students")
def get_students(
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] != "HEAD":

        raise HTTPException(
            status_code=403,
            detail="Only Head can view all students"
        )


    workbook = get_workbook()

    sheet = workbook[
        "Students"
    ]


    result = []


    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        result.append({

            "student_id":
                row[0],

            "name":
                row[1],

            "roll_no":
                row[2],

            "section":
                row[3],

            "branch":
                row[4],

            "email":
                row[5],

            "registered_date":
                row[8]

        })


    workbook.close()

    return result


# =========================================================
# GET QR
# =========================================================

@app.get("/qr/{student_id}")
def get_qr(
    student_id: str
):

    qr_path = (
        QR_FOLDER /
        f"{student_id}.png"
    )


    if not qr_path.exists():

        raise HTTPException(
            status_code=404,
            detail="QR code not found"
        )


    return FileResponse(
        qr_path,
        media_type="image/png"
    )


# =========================================================
# CREATE SESSION
# HEAD + SESSION CREATOR
# =========================================================

@app.post("/sessions")
def create_session(
    data: SessionCreate,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] not in [
        "HEAD",
        "SESSION_CREATOR"
    ]:

        raise HTTPException(
            status_code=403,
            detail="You do not have permission to create a session"
        )


    event = data.event.strip()


    if not event:

        raise HTTPException(
            status_code=400,
            detail="Meeting name is required"
        )


    workbook = get_workbook()

    sheet = workbook[
        "Sessions"
    ]


    now = datetime.now()


    session_id = (
        "S-"
        + now.strftime(
            "%Y%m%d%H%M%S"
        )
        + "-"
        + secrets.token_hex(
            3
        ).upper()
    )


    sheet.append([

        session_id,

        event,

        now.strftime(
            "%Y-%m-%d"
        ),

        now.strftime(
            "%H:%M:%S"
        ),

        current_user["user_id"]

    ])


    workbook.save(
        EXCEL_FILE
    )

    workbook.close()


    # Create separate Excel for this meeting
    create_session_excel(
        session_id,
        event
    )


    return {

        "message":
            "Session created successfully",

        "session_id":
            session_id,

        "event":
            event

    }


# =========================================================
# GET SESSIONS
# =========================================================

@app.get("/sessions")
def get_sessions(
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] not in [
        "HEAD",
        "SESSION_CREATOR"
    ]:

        raise HTTPException(
            status_code=403,
            detail="You do not have permission to view sessions"
        )


    workbook = get_workbook()

    sheet = workbook[
        "Sessions"
    ]


    result = []


    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        # Session Creator sees only sessions
        # created by themselves

        if (
            current_user["role"]
            == "SESSION_CREATOR"
            and
            str(row[4])
            != str(current_user["user_id"])
        ):

            continue


        result.append({

            "session_id":
                row[0],

            "event":
                row[1],

            "date":
                row[2],

            "start_time":
                row[3],

            "created_by":
                row[4]

        })


    workbook.close()

    return result


# =========================================================
# MARK ATTENDANCE
# =========================================================

@app.post("/attendance")
def mark_attendance(
    data: AttendanceRequest,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] not in [
        "HEAD",
        "SESSION_CREATOR"
    ]:

        raise HTTPException(
            status_code=403,
            detail="Only authorized attendance staff can scan QR codes"
        )


    workbook = get_workbook()


    # =====================================================
    # FIND SESSION
    # =====================================================

    sessions = workbook[
        "Sessions"
    ]

    session = None


    for row in sessions.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == str(
            data.session_id
        ):

            session = {

                "session_id":
                    row[0],

                "event":
                    row[1],

                "date":
                    row[2],

                "start_time":
                    row[3],

                "created_by":
                    row[4]

            }

            break


    if not session:

        workbook.close()

        raise HTTPException(
            status_code=404,
            detail="Session not found"
        )


    # Session Creator can scan
    # only their own session

    if (
        current_user["role"]
        == "SESSION_CREATOR"
        and
        str(session["created_by"])
        != str(current_user["user_id"])
    ):

        workbook.close()

        raise HTTPException(
            status_code=403,
            detail="You can only scan attendance for your own session"
        )


    # =====================================================
    # FIND STUDENT USING QR TOKEN
    # =====================================================

    students = workbook[
        "Students"
    ]

    student = None


    for row in students.iter_rows(
        min_row=2,
        values_only=True
    ):

        # QR Token is column 8
        if str(row[7]) == str(
            data.qr_token
        ):

            student = {

                "student_id":
                    row[0],

                "name":
                    row[1],

                "roll_no":
                    row[2],

                "section":
                    row[3],

                "branch":
                    row[4],

                "email":
                    row[5]

            }

            break


    if not student:

        workbook.close()

        raise HTTPException(
            status_code=404,
            detail="Invalid student QR code"
        )


    # =====================================================
    # DUPLICATE CHECK
    # =====================================================

    attendance = workbook[
        "Attendance"
    ]


    for row in attendance.iter_rows(
        min_row=2,
        values_only=True
    ):

        if (
            str(row[0])
            == str(data.session_id)
            and
            str(row[2])
            == str(student["student_id"])
        ):

            workbook.close()

            raise HTTPException(
                status_code=400,
                detail="Student already marked present"
            )


    # =====================================================
    # MARK ATTENDANCE IN MASTER EXCEL
    # =====================================================

    now = datetime.now()


    attendance.append([

        data.session_id,

        session["event"],

        student["student_id"],

        student["name"],

        student["roll_no"],

        student["section"],

        student["branch"],

        student["email"],

        now.strftime(
            "%Y-%m-%d"
        ),

        now.strftime(
            "%H:%M:%S"
        ),

        "Present"

    ])


    workbook.save(
        EXCEL_FILE
    )

    workbook.close()


    # =====================================================
    # ALSO SAVE INTO THIS MEETING'S EXCEL
    # =====================================================

    append_to_session_excel(

        session["session_id"],

        session["event"],

        student,

        now

    )


    return {

        "message":
            "Attendance marked successfully",

        "student_id":
            student["student_id"],

        "name":
            student["name"],

        "status":
            "Present",

        "time":
            now.strftime(
                "%H:%M:%S"
            )

    }


# =========================================================
# ATTENDANCE REPORT
# =========================================================

@app.get("/attendance/{session_id}")
def get_attendance(
    session_id: str,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] not in [
        "HEAD",
        "SESSION_CREATOR"
    ]:

        raise HTTPException(
            status_code=403,
            detail="You do not have permission to view this report"
        )


    workbook = get_workbook()


    # =====================================================
    # FIND SESSION
    # =====================================================

    sessions = workbook[
        "Sessions"
    ]

    session = None


    for row in sessions.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == str(
            session_id
        ):

            session = {

                "session_id":
                    row[0],

                "event":
                    row[1],

                "created_by":
                    row[4]

            }

            break


    if not session:

        workbook.close()

        raise HTTPException(
            status_code=404,
            detail="Session not found"
        )


    # Session Creator can view
    # only own session

    if (
        current_user["role"]
        == "SESSION_CREATOR"
        and
        str(session["created_by"])
        != str(current_user["user_id"])
    ):

        workbook.close()

        raise HTTPException(
            status_code=403,
            detail="You can only view your own session"
        )


    students = workbook[
        "Students"
    ]

    attendance = workbook[
        "Attendance"
    ]


    total_students = max(
        students.max_row - 1,
        0
    )


    records = []


    for row in attendance.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == str(
            session_id
        ):

            records.append({

                "student_id":
                    row[2],

                "name":
                    row[3],

                "roll_no":
                    row[4],

                "section":
                    row[5],

                "branch":
                    row[6],

                "email":
                    row[7],

                "date":
                    row[8],

                "time":
                    row[9],

                "status":
                    row[10]

            })


    present = len(records)


    workbook.close()


    return {

        "session_id":
            session_id,

        "event":
            session["event"],

        "total_students":
            total_students,

        "present":
            present,

        "absent":
            max(
                total_students - present,
                0
            ),

        "attendance":
            records

    }


# =========================================================
# DOWNLOAD PARTICULAR SESSION EXCEL
# =========================================================

@app.get("/download/{session_id}")
def download_session_excel(
    session_id: str,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] not in [
        "HEAD",
        "SESSION_CREATOR"
    ]:

        raise HTTPException(
            status_code=403,
            detail="You do not have permission to download reports"
        )


    workbook = get_workbook()

    sessions = workbook[
        "Sessions"
    ]

    session = None


    for row in sessions.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == str(
            session_id
        ):

            session = {

                "session_id":
                    row[0],

                "event":
                    row[1],

                "created_by":
                    row[4]

            }

            break


    if not session:

        workbook.close()

        raise HTTPException(
            status_code=404,
            detail="Session not found"
        )


    # Session Creator can download
    # only own session

    if (
        current_user["role"]
        == "SESSION_CREATOR"
        and
        str(session["created_by"])
        != str(current_user["user_id"])
    ):

        workbook.close()

        raise HTTPException(
            status_code=403,
            detail="You can only download your own session"
        )


    workbook.close()


    path = session_excel_path(

        session["session_id"],

        session["event"]

    )


    # Create if somehow missing

    if not path.exists():

        create_session_excel(

            session["session_id"],

            session["event"]

        )


    download_name = (
        f"{safe_filename(session['event'])}.xlsx"
    )


    return FileResponse(

        path,

        filename=download_name,

        media_type=
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )


# =========================================================
# DELETE SESSION
# HEAD ONLY
# =========================================================

@app.delete("/sessions/{session_id}")
def delete_session(
    session_id: str,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    # ONLY HEAD CAN DELETE

    if current_user["role"] != "HEAD":

        raise HTTPException(
            status_code=403,
            detail="Only Head can delete a meeting"
        )


    workbook = get_workbook()


    sessions = workbook[
        "Sessions"
    ]


    session = None


    # =====================================================
    # FIND SESSION
    # =====================================================

    for row in sessions.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[0]) == str(
            session_id
        ):

            session = {

                "session_id":
                    row[0],

                "event":
                    row[1]

            }

            break


    if not session:

        workbook.close()

        raise HTTPException(
            status_code=404,
            detail="Session not found"
        )


    # =====================================================
    # DELETE SESSION ROW
    # =====================================================

    delete_row_number = None


    for row_number in range(
        2,
        sessions.max_row + 1
    ):

        if str(
            sessions.cell(
                row_number,
                1
            ).value
        ) == str(session_id):

            delete_row_number = row_number

            break


    if delete_row_number:

        sessions.delete_rows(
            delete_row_number,
            1
        )


    # =====================================================
    # DELETE ATTENDANCE RECORDS
    # =====================================================

    attendance = workbook[
        "Attendance"
    ]


    rows_to_delete = []


    for row_number in range(
        2,
        attendance.max_row + 1
    ):

        if str(
            attendance.cell(
                row_number,
                1
            ).value
        ) == str(session_id):

            rows_to_delete.append(
                row_number
            )


    for row_number in reversed(
        rows_to_delete
    ):

        attendance.delete_rows(
            row_number,
            1
        )


    workbook.save(
        EXCEL_FILE
    )

    workbook.close()


    # =====================================================
    # DELETE SESSION EXCEL FILE
    # =====================================================

    session_file = session_excel_path(

        session["session_id"],

        session["event"]

    )


    if session_file.exists():

        try:

            session_file.unlink()

        except Exception:

            pass


    return {

        "success": True,

        "message":
            "Meeting and attendance data deleted successfully"

    }


# =========================================================
# HEAD CREATES SESSION CREATOR
# =========================================================

@app.post("/session-creators")
def create_session_creator(
    data: SessionCreatorCreate,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] != "HEAD":

        raise HTTPException(
            status_code=403,
            detail="Only Head can give session creation access"
        )


    workbook = get_workbook()


    if find_user(
        workbook,
        data.user_id
    ):

        workbook.close()

        raise HTTPException(
            status_code=400,
            detail="User ID already exists"
        )


    users = workbook[
        "Users"
    ]


    users.append([

        data.user_id,

        data.name,

        data.email,

        hash_password(
            data.password
        ),

        "SESSION_CREATOR",

        "YES",

        datetime.now().strftime(
            "%Y-%m-%d"
        )

    ])


    workbook.save(
        EXCEL_FILE
    )

    workbook.close()


    return {

        "message":
            "Session Creator account created",

        "user_id":
            data.user_id,

        "role":
            "SESSION_CREATOR"

    }


# =========================================================
# VIEW SESSION CREATORS
# HEAD ONLY
# =========================================================

@app.get("/session-creators")
def get_session_creators(
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] != "HEAD":

        raise HTTPException(
            status_code=403,
            detail="Only Head can view session creators"
        )


    workbook = get_workbook()

    users = workbook[
        "Users"
    ]


    result = []


    for row in users.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[4] == "SESSION_CREATOR":

            result.append({

                "user_id":
                    row[0],

                "name":
                    row[1],

                "email":
                    row[2],

                "active":
                    row[5]

            })


    workbook.close()

    return result


# =========================================================
# DISABLE SESSION CREATOR
# =========================================================

@app.delete("/session-creators/{user_id}")
def disable_session_creator(
    user_id: str,
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] != "HEAD":

        raise HTTPException(
            status_code=403,
            detail="Only Head can remove access"
        )


    workbook = get_workbook()

    users = workbook[
        "Users"
    ]


    found = False


    for row in users.iter_rows(
        min_row=2
    ):

        if row[0].value == user_id:

            if row[4].value != "SESSION_CREATOR":

                workbook.close()

                raise HTTPException(
                    status_code=400,
                    detail="This is not a Session Creator account"
                )


            row[5].value = "NO"

            found = True

            break


    if not found:

        workbook.close()

        raise HTTPException(
            status_code=404,
            detail="Session Creator not found"
        )


    workbook.save(
        EXCEL_FILE
    )

    workbook.close()


    # Remove active login tokens
    tokens_to_remove = []

    for token, user in ACTIVE_TOKENS.items():

        if str(
            user["user_id"]
        ) == str(user_id):

            tokens_to_remove.append(
                token
            )


    for token in tokens_to_remove:

        ACTIVE_TOKENS.pop(
            token,
            None
        )


    return {

        "message":
            "Session creation access removed"

    }


# =========================================================
# STUDENT OWN PROFILE
# =========================================================
@app.get("/my-profile")
def my_profile(
    authorization: str | None = Header(None)
):
    current_user = get_current_user(
        authorization
    )

    if current_user["role"] != "STUDENT":
        raise HTTPException(
            status_code=403,
            detail="Student access only"
        )

    workbook = get_workbook()

    students = workbook["Students"]

    student = None

    for row in students.iter_rows(
        min_row=2,
        values_only=True
    ):
        if str(row[0]) == str(
            current_user["user_id"]
        ):
            student = {
                "student_id": row[0],
                "name": row[1],
                "roll_no": row[2],
                "section": row[3],
                "branch": row[4],
                "email": row[5],
                "registered_date": row[8]
            }
            break

    workbook.close()

    if not student:
        raise HTTPException(
            status_code=404,
            detail="Student profile not found"
        )

    return student


# =========================================================
# STUDENT OWN QR CODE
# =========================================================
@app.get("/my-qr")
def my_qr(
    authorization: str | None = Header(None)
):
    current_user = get_current_user(
        authorization
    )

    if current_user["role"] != "STUDENT":
        raise HTTPException(
            status_code=403,
            detail="Student access only"
        )

    student_id = current_user["user_id"]

    qr_path = (
        QR_FOLDER /
        f"{student_id}.png"
    )

    if not qr_path.exists():
        raise HTTPException(
            status_code=404,
            detail="QR code not found"
        )

    return FileResponse(
        qr_path,
        media_type="image/png"
    )
# =========================================================
# STUDENT OWN ATTENDANCE
# =========================================================

@app.get("/my-attendance")
def my_attendance(
    authorization: str | None = Header(None)
):

    current_user = get_current_user(
        authorization
    )


    if current_user["role"] != "STUDENT":

        raise HTTPException(
            status_code=403,
            detail="Student access only"
        )


    workbook = get_workbook()

    attendance = workbook[
        "Attendance"
    ]


    records = []


    for row in attendance.iter_rows(
        min_row=2,
        values_only=True
    ):

        if str(row[2]) == str(
            current_user["user_id"]
        ):

            records.append({

                "session_id":
                    row[0],

                "event":
                    row[1],

                "name":
                    row[3],

                "roll_no":
                    row[4],

                "section":
                    row[5],

                "branch":
                    row[6],

                "email":
                    row[7],

                "date":
                    row[8],

                "time":
                    row[9],

                "status":
                    row[10]

            })


    workbook.close()

    return records


# =========================================================
# HEALTH
# =========================================================

@app.get("/health")
def health():

    return {
        "status":
            "running"
    }


# =========================================================
# HOME
# =========================================================

@app.get("/")
def home():

    return {

        "message":
            "IEEE CIS Attendance System API is running"

    }